import os
import json
import csv
import io
from datetime import datetime, timedelta
from io import BytesIO
import smtplib
from email.message import EmailMessage
import openpyxl
from PIL import Image

from fastapi import FastAPI, Request, Form, UploadFile, File, Response
from fastapi.responses import HTMLResponse, FileResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles

import firebase_admin
from firebase_admin import credentials, firestore, storage

# ==========================================
# FIREBASE BULUT AYARLARI
# ==========================================
firebase_env = os.getenv("FIREBASE_JSON")
if firebase_env:
    cred_dict = json.loads(firebase_env)
    cred = credentials.Certificate(cred_dict)
else:
    cred = credentials.Certificate("firebase-key.json")

firebase_admin.initialize_app(cred, {
    'storageBucket': 'kbis-okul.firebasestorage.app'
})

db = firestore.client()
bucket = storage.bucket()
# ==========================================

app = FastAPI(title="K-BİS")
templates = Jinja2Templates(directory="templates")

os.makedirs("static", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

def mail_gonder(alici_eposta, konu, icerik):
    GONDERICI_MAIL = os.getenv("MAIL_ADRES", "okul.kbis@gmail.com") 
    GONDERICI_SIFRE = os.getenv("MAIL_SIFRE", "uygulama_sifresi_buraya") 
    try:
        msg = EmailMessage()
        msg.set_content(icerik, subtype='html')
        msg['Subject'] = konu
        msg['From'] = GONDERICI_MAIL
        msg['To'] = alici_eposta
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(GONDERICI_MAIL, GONDERICI_SIFRE)
        server.send_message(msg)
        server.quit()
    except Exception as e:
        print(f"Mail Hatası: {e}")

def format_bildirimler(kayitlar):
    liste = []
    for d in kayitlar:
        ac = d.get("aciliyet", 1)
        if ac > 3: ac = 3
        if d.get("tur") == "Risk":
            d["seviye_metni"] = {1: "Az Riskli", 2: "Riskli", 3: "Çok Riskli"}.get(ac, "Bilinmiyor")
        else:
            d["seviye_metni"] = {1: "Normal", 2: "Acil", 3: "Çok Acil"}.get(ac, "Bilinmiyor")
        d["seviye_renk"] = {1: "#10b981", 2: "#f59e0b", 3: "#ef4444"}.get(ac, "#64748b")
        liste.append(d)
    return liste

def idareci_verilerini_getir(idareci_id, mesaj=""):
    kullanici_doc = db.collection("Kullanicilar").document(idareci_id).get()
    kullanici = kullanici_doc.to_dict()
    kullanici["id"] = kullanici_doc.id

    tum_kullanicilar_ref = db.collection("Kullanicilar").get()
    kullanicilar_dict = {doc.id: doc.to_dict() for doc in tum_kullanicilar_ref}
    
    tum_mekanlar_ref = db.collection("Mekanlar").get()
    mekanlar_dict = {doc.id: doc.to_dict() for doc in tum_mekanlar_ref}

    sorumlular = [{"id": k_id, **k_data} for k_id, k_data in kullanicilar_dict.items() if k_data.get("rol") in ['Şef', 'Müdür Yardımcısı', 'Alan Şefi']]
    sorumlular = sorted(sorumlular, key=lambda x: (x.get('rol', ''), x.get('ad_soyad', '')))

    tum_k_liste = [{"id": k_id, **k_data} for k_id, k_data in kullanicilar_dict.items()]
    tum_k_liste = sorted(tum_k_liste, key=lambda x: (x.get('rol', ''), x.get('ad_soyad', '')))

    tum_m_liste = [{"id": m_id, **m_data} for m_id, m_data in mekanlar_dict.items()]
    tum_m_liste = sorted(tum_m_liste, key=lambda x: (x.get('bina_adi', ''), x.get('kat', ''), x.get('alan_adi', '')))

    bildirimler_ref = db.collection("Bildirimler").order_by("aciliyet", direction=firestore.Query.DESCENDING).order_by("tarih", direction=firestore.Query.DESCENDING).get()
    
    ham_bildirimler = []
    for doc in bildirimler_ref:
        b_data = doc.to_dict()
        b_data["id"] = doc.id
        
        mekan = mekanlar_dict.get(b_data.get("mekan_id"), {})
        b_data["bina_adi"] = mekan.get("bina_adi", "Bilinmiyor")
        b_data["kat"] = mekan.get("kat", "")
        b_data["alan_adi"] = mekan.get("alan_adi", "")
        
        bildiren = kullanicilar_dict.get(b_data.get("bildiren_id"), {})
        b_data["bildiren_kisi"] = bildiren.get("ad_soyad", "Bilinmiyor")
        
        sef = kullanicilar_dict.get(b_data.get("atanan_sef_id"), {})
        b_data["sef_adi"] = sef.get("ad_soyad", "")
        
        ham_bildirimler.append(b_data)

    bildirimler = format_bildirimler(ham_bildirimler)
        
    return {"kullanici": kullanici, "bildirimler": bildirimler, "sorumlular": sorumlular, "tum_kullanicilar": tum_k_liste, "tum_mekanlar": tum_m_liste, "mesaj": mesaj}

def ogretmen_verilerini_getir(kullanici_id, mesaj=""):
    kullanici_doc = db.collection("Kullanicilar").document(kullanici_id).get()
    kullanici = kullanici_doc.to_dict()
    kullanici["id"] = kullanici_doc.id

    tum_mekanlar_ref = db.collection("Mekanlar").get()
    m_dict = {doc.id: doc.to_dict() for doc in tum_mekanlar_ref}

    bildirimler_ref = db.collection("Bildirimler").where("bildiren_id", "==", kullanici_id).get()
    
    ham_bildirimler = []
    for doc in bildirimler_ref:
        b_data = doc.to_dict()
        b_data["id"] = doc.id
        mekan = m_dict.get(b_data.get("mekan_id"), {})
        b_data["bina_adi"] = mekan.get("bina_adi", "Bilinmiyor")
        b_data["kat"] = mekan.get("kat", "")
        b_data["alan_adi"] = mekan.get("alan_adi", "")
        ham_bildirimler.append(b_data)

    ham_bildirimler = sorted(ham_bildirimler, key=lambda x: x.get('tarih', ''), reverse=True)
    gecmis_bildirimler = format_bildirimler(ham_bildirimler)
        
    return {"kullanici": kullanici, "gecmis_bildirimler": gecmis_bildirimler, "mesaj": mesaj}

@app.get("/kurulum", response_class=HTMLResponse)
async def veritabani_kurulum():
    yonetici = db.collection("Kullanicilar").where("eposta", "==", "admin@kucukyali.edu.tr").get()
    if not yonetici:
        db.collection("Kullanicilar").add({
            "eposta": "admin@kucukyali.edu.tr",
            "sifre": "123456",
            "ad_soyad": "Sistem Yöneticisi",
            "rol": "Müdür"
        })
        return "✅ Firebase kurulumu başarılı. İlk yönetici hesabı oluşturuldu. Ana sayfadan giriş yapabilirsiniz."
    return "Yönetici zaten mevcut."

@app.get("/", response_class=HTMLResponse)
async def login_sayfasi(request: Request):
    return templates.TemplateResponse(request=request, name="login.html", context={"mesaj": ""})

@app.post("/giris")
async def giris_yap(request: Request, eposta: str = Form(...), sifre: str = Form(...)):
    try:
        users = db.collection("Kullanicilar").where("eposta", "==", eposta).where("sifre", "==", sifre).get()
        
        if users:
            kullanici_doc = users[0]
            kullanici = kullanici_doc.to_dict()
            kullanici["id"] = kullanici_doc.id
            rol = kullanici.get("rol")
            
            if rol in ["Müdür", "Teknik Müdür Yardımcısı", "İdareci"]:
                context = idareci_verilerini_getir(kullanici["id"])
                return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)
            elif rol == "Öğretmen":
                context = ogretmen_verilerini_getir(kullanici["id"])
                return templates.TemplateResponse(request=request, name="ogretmen_panel.html", context=context)
            elif rol in ["Şef", "Müdür Yardımcısı", "Alan Şefi"]:
                isler_ref = db.collection("Bildirimler").where("atanan_sef_id", "==", kullanici["id"]).where("durum", "in", ["İşlemde (Şefte)", "İşlem Bekliyor"]).get()
                
                tum_kullanicilar_ref = db.collection("Kullanicilar").get()
                k_dict = {doc.id: doc.to_dict() for doc in tum_kullanicilar_ref}
                tum_mekanlar_ref = db.collection("Mekanlar").get()
                m_dict = {doc.id: doc.to_dict() for doc in tum_mekanlar_ref}

                ham_isler = []
                for doc in isler_ref:
                    is_data = doc.to_dict()
                    is_data["id"] = doc.id
                    mekan = m_dict.get(is_data.get("mekan_id"), {})
                    is_data["bina_adi"] = mekan.get("bina_adi", "Bilinmiyor")
                    is_data["kat"] = mekan.get("kat", "")
                    is_data["alan_adi"] = mekan.get("alan_adi", "")
                    bildiren = k_dict.get(is_data.get("bildiren_id"), {})
                    is_data["bildiren_kisi"] = bildiren.get("ad_soyad", "Bilinmiyor")
                    ham_isler.append(is_data)
                
                ham_isler = sorted(ham_isler, key=lambda x: (x.get('aciliyet', 1), x.get('tarih', '')), reverse=True)
                atanan_isler = format_bildirimler(ham_isler)
                return templates.TemplateResponse(request=request, name="sef_panel.html", context={"kullanici": dict(kullanici), "isler": atanan_isler})
                
        return templates.TemplateResponse(request=request, name="login.html", context={"mesaj": "Hatalı e-posta veya şifre.", "basarili": False})
    except Exception as e:
        import traceback
        return HTMLResponse(content=f"<h1 style='color:red;'>SİSTEM HATASI:</h1><pre>{traceback.format_exc()}</pre>", status_code=500)

@app.get("/idareci_dashboard/{idareci_id}")
async def idareci_dashboard(request: Request, idareci_id: str):
    context = idareci_verilerini_getir(idareci_id)
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.post("/sifre_degistir")
async def sifre_degistir(request: Request, kullanici_id: str = Form(...), eski_sifre: str = Form(...), yeni_sifre: str = Form(...)):
    user_ref = db.collection("Kullanicilar").document(kullanici_id)
    user_doc = user_ref.get()
    
    if user_doc.exists:
        if user_doc.to_dict().get("sifre") == eski_sifre:
            user_ref.update({"sifre": yeni_sifre})
            mesaj = "✅ Şifreniz başarıyla güncellendi."
        else:
            mesaj = "❌ HATA: Eski şifrenizi yanlış girdiniz!"
    else:
        mesaj = "❌ Sistem Hatası: Kullanıcı bulunamadı."
        
    context = ogretmen_verilerini_getir(kullanici_id, mesaj)
    return templates.TemplateResponse(request=request, name="ogretmen_panel.html", context=context)

@app.post("/kullanici_ekle")
async def kullanici_ekle(request: Request, idareci_id: str = Form(...), ad_soyad: str = Form(...), yeni_eposta: str = Form(...), yeni_sifre: str = Form(...), rol: str = Form(...)):
    istek_yapan = db.collection("Kullanicilar").document(idareci_id).get()
    if not istek_yapan.exists or istek_yapan.to_dict().get("rol") not in ["Müdür", "Teknik Müdür Yardımcısı", "İdareci"]:
        return templates.TemplateResponse(request=request, name="login.html", context={"mesaj": "❌ YETKİSİZ İŞLEM!"})
    
    mevcut = db.collection("Kullanicilar").where("eposta", "==", yeni_eposta).get()
    if mevcut:
        mesaj = "❌ HATA: E-posta adresi kullanılıyor!"
    else:
        db.collection("Kullanicilar").add({
            "eposta": yeni_eposta,
            "sifre": yeni_sifre,
            "ad_soyad": ad_soyad,
            "rol": rol
        })
        mesaj = f"✅ {ad_soyad}, sisteme {rol} olarak eklendi."
        
    context = idareci_verilerini_getir(idareci_id, mesaj)
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.post("/kullanici_sil")
async def kullanici_sil(request: Request, idareci_id: str = Form(...), silinecek_id: str = Form(...)):
    db.collection("Kullanicilar").document(silinecek_id).delete()
    context = idareci_verilerini_getir(idareci_id, "🗑️ Kullanıcı sistemden silindi.")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.get("/ornek_kullanici_excel")
async def ornek_kullanici_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ad Soyad", "E-Posta", "Şifre", "Rol"])
    ws.append(["Ahmet Yılmaz", "ahmet@okul.com", "123456", "Öğretmen"])
    ws.append(["Mehmet Demir", "mehmet@okul.com", "123456", "Şef"])
    stream = BytesIO()
    wb.save(stream)
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=ornek_kullanicilar.xlsx"})

@app.post("/toplu_kullanici_yukle")
async def toplu_kullanici_yukle(request: Request, idareci_id: str = Form(...), excel_dosyasi: UploadFile = File(...)):
    contents = await excel_dosyasi.read()
    wb = openpyxl.load_workbook(filename=BytesIO(contents))
    ws = wb.active
    
    eklenen = 0
    mevcut_epostalar = [doc.to_dict().get("eposta") for doc in db.collection("Kullanicilar").stream()]
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]: continue
        ad_soyad, eposta, sifre, rol = str(row[0]), str(row[1]), str(row[2]), str(row[3])
        
        if eposta not in mevcut_epostalar:
            db.collection("Kullanicilar").add({
                "ad_soyad": ad_soyad,
                "eposta": eposta,
                "sifre": sifre,
                "rol": rol
            })
            mevcut_epostalar.append(eposta)
            eklenen += 1
            
    context = idareci_verilerini_getir(idareci_id, f"✅ Excel'den {eklenen} adet yeni kullanıcı başarıyla sisteme eklendi.")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)


@app.post("/mekan_kaydet")
async def mekan_kaydet(request: Request, bina_adi: str = Form(...), kat: str = Form(...), alan_adi: str = Form(...), idareci_id: str = Form(...)):
    db.collection("Mekanlar").add({
        "bina_adi": bina_adi,
        "kat": kat,
        "alan_adi": alan_adi
    })
    context = idareci_verilerini_getir(idareci_id, f"✅ {bina_adi} eklendi!")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.post("/mekan_sil")
async def mekan_sil(request: Request, idareci_id: str = Form(...), silinecek_id: str = Form(...)):
    db.collection("Mekanlar").document(silinecek_id).delete()
    context = idareci_verilerini_getir(idareci_id, "🗑️ Seçilen mekan/konum başarıyla silindi.")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.get("/ornek_mekan_excel")
async def ornek_mekan_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Bina Adı", "Kat", "Alan Adı"])
    ws.append(["A Blok", "Zemin Kat", "Öğretmenler Odası"])
    stream = BytesIO()
    wb.save(stream)
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=ornek_mekanlar.xlsx"})

@app.post("/toplu_mekan_yukle")
async def toplu_mekan_yukle(request: Request, idareci_id: str = Form(...), excel_dosyasi: UploadFile = File(...)):
    contents = await excel_dosyasi.read()
    wb = openpyxl.load_workbook(filename=BytesIO(contents))
    ws = wb.active
    
    eklenen = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]: continue
        bina_adi, kat, alan_adi = str(row[0]), str(row[1]), str(row[2])
        
        db.collection("Mekanlar").add({
            "bina_adi": bina_adi,
            "kat": kat,
            "alan_adi": alan_adi
        })
        eklenen += 1
            
    context = idareci_verilerini_getir(idareci_id, f"✅ Excel'den {eklenen} adet yeni mekan başarıyla sisteme eklendi.")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)


@app.post("/sefe_ata")
async def sefe_ata(request: Request, bildirim_id: str = Form(...), sef_id: str = Form(...), idareci_id: str = Form(...)):
    db.collection("Bildirimler").document(bildirim_id).update({
        "durum": "İşlemde (Şefte)",
        "atanan_sef_id": sef_id,
        "atama_tarihi": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
    
    atanan_personel_doc = db.collection("Kullanicilar").document(sef_id).get()
    if atanan_personel_doc.exists:
        personel = atanan_personel_doc.to_dict()
        bildirim_doc = db.collection("Bildirimler").document(bildirim_id).get()
        if bildirim_doc.exists:
            bildirim = bildirim_doc.to_dict()
            mekan_doc = db.collection("Mekanlar").document(bildirim.get("mekan_id", "")).get()
            mekan_adi = f"{mekan_doc.to_dict().get('bina_adi')} - {mekan_doc.to_dict().get('alan_adi')}" if mekan_doc.exists else "Bilinmeyen Konum"
            
            konu = f"📋 Yeni İş Emri Atandı: {bildirim.get('tur')} Bildirimi"
            icerik = f"<h3 style='color:#3b82f6;'>Sayın {personel.get('ad_soyad')},</h3><p>Sisteme girilen yeni bir arıza/risk bildirimi idare tarafından tarafınıza atanmıştır.</p><ul><li><b>Konum:</b> {mekan_adi}</li><li><b>Tür:</b> {bildirim.get('tur')}</li><li><b>Açıklama:</b> {bildirim.get('aciklama')}</li></ul><p>Lütfen K-BİS paneline giriş yaparak iş emri detaylarını inceleyiniz.</p>"
            mail_gonder(personel.get("eposta"), konu, icerik)

    context = idareci_verilerini_getir(idareci_id, "✅ İş emri atandı ve ilgili personele bilgi maili gönderildi!")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.post("/personel_ata")
async def personel_ata(request: Request, bildirim_id: str = Form(...), gorevlendirilen_personel: str = Form(...), sef_id: str = Form(...)):
    db.collection("Bildirimler").document(bildirim_id).update({
        "gorevlendirilen_personel": gorevlendirilen_personel
    })
    
    isler_ref = db.collection("Bildirimler").where("atanan_sef_id", "==", sef_id).where("durum", "in", ["İşlemde (Şefte)", "İşlem Bekliyor"]).get()
    tum_kullanicilar_ref = db.collection("Kullanicilar").get()
    k_dict = {doc.id: doc.to_dict() for doc in tum_kullanicilar_ref}
    tum_mekanlar_ref = db.collection("Mekanlar").get()
    m_dict = {doc.id: doc.to_dict() for doc in tum_mekanlar_ref}

    ham_isler = []
    for doc in isler_ref:
        is_data = doc.to_dict()
        is_data["id"] = doc.id
        mekan = m_dict.get(is_data.get("mekan_id"), {})
        is_data["bina_adi"] = mekan.get("bina_adi", "Bilinmiyor")
        is_data["kat"] = mekan.get("kat", "")
        is_data["alan_adi"] = mekan.get("alan_adi", "")
        bildiren = k_dict.get(is_data.get("bildiren_id"), {})
        is_data["bildiren_kisi"] = bildiren.get("ad_soyad", "Bilinmiyor")
        ham_isler.append(is_data)
    
    ham_isler = sorted(ham_isler, key=lambda x: (x.get('aciliyet', 1), x.get('tarih', '')), reverse=True)
    atanan_isler = format_bildirimler(ham_isler)
    
    kullanici = db.collection("Kullanicilar").document(sef_id).get().to_dict()
    kullanici["id"] = sef_id
    
    return templates.TemplateResponse(request=request, name="sef_panel.html", context={"kullanici": kullanici, "isler": atanan_isler, "mesaj": "✅ Personel görevlendirmesi kaydedildi!"})

@app.get("/is_emri_yazdir/{bildirim_id}", response_class=HTMLResponse)
async def is_emri_yazdir(request: Request, bildirim_id: str):
    doc = db.collection("Bildirimler").document(bildirim_id).get()
    if not doc.exists:
        return HTMLResponse("Bildirim bulunamadı.")
    
    b = doc.to_dict()
    b["id"] = doc.id
    
    mekan = db.collection("Mekanlar").document(b.get("mekan_id", "")).get().to_dict() or {}
    b["bina_adi"] = mekan.get("bina_adi", "")
    b["kat"] = mekan.get("kat", "")
    b["alan_adi"] = mekan.get("alan_adi", "")
    
    bildiren = db.collection("Kullanicilar").document(b.get("bildiren_id", "")).get().to_dict() or {}
    b["bildiren_kisi"] = bildiren.get("ad_soyad", "")
    
    if b.get("atanan_sef_id"):
        sef = db.collection("Kullanicilar").document(b.get("atanan_sef_id", "")).get().to_dict() or {}
        b["sef_adi"] = sef.get("ad_soyad", "")
    else:
        b["sef_adi"] = ""
        
    return templates.TemplateResponse(request=request, name="is_emri_yazdir.html", context={"islem": b})

@app.post("/idareci_kapat")
async def idareci_kapat(request: Request, bildirim_id: str = Form(...), idareci_id: str = Form(...)):
    db.collection("Bildirimler").document(bildirim_id).update({
        "durum": "Kapatıldı",
        "kapanis_tarihi": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
    
    bildirim_doc = db.collection("Bildirimler").document(bildirim_id).get()
    bildirim = bildirim_doc.to_dict()
    ogretmen_doc = db.collection("Kullanicilar").document(bildirim.get("bildiren_id", "")).get()
    
    if ogretmen_doc.exists:
        ogretmen = ogretmen_doc.to_dict()
        konu = "✅ K-BİS: Bildirdiğiniz Sorun Çözüldü"
        icerik = f"<h3 style='color:#10b981;'>Merhaba {ogretmen.get('ad_soyad')},</h3><p>Sisteme girdiğiniz bildirim Sorumlu Personel tarafından çözülmüş ve İdare tarafından onaylanarak kapatılmıştır.</p><div style='background-color: #f9f9f9; padding: 15px; border-left: 4px solid #10b981; margin: 15px 0;'><p><b>Çözüm Notu:</b><br><i>{bildirim.get('sef_cozum_notu', '')}</i></p></div>"
        if bildirim.get('cozum_fotografi'):
            icerik += f"<p><b>Onarım Sonrası Fotoğrafı:</b></p><img src='{bildirim.get('cozum_fotografi')}' alt='Çözüm Görseli' style='max-width: 400px; border-radius: 8px; border: 1px solid #ccc;'>"
        mail_gonder(ogretmen.get("eposta"), konu, icerik)
            
    context = idareci_verilerini_getir(idareci_id, "✅ İş kapatıldı ve bildirim sahibine mail gönderildi!")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.post("/bildirim_iptal")
async def bildirim_iptal(request: Request, bildirim_id: str = Form(...), iptal_notu: str = Form(...), idareci_id: str = Form(...)):
    db.collection("Bildirimler").document(bildirim_id).update({
        "durum": "Yanlış Bildirim",
        "sef_cozum_notu": f"İPTAL EDİLDİ: {iptal_notu}",
        "kapanis_tarihi": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
    
    bildirim_doc = db.collection("Bildirimler").document(bildirim_id).get()
    bildirim = bildirim_doc.to_dict()
    ogretmen_doc = db.collection("Kullanicilar").document(bildirim.get("bildiren_id", "")).get()
    
    if ogretmen_doc.exists:
        ogretmen = ogretmen_doc.to_dict()
        konu = "❌ K-BİS: Bildiriminiz İptal Edildi"
        icerik = f"<h3>Merhaba {ogretmen.get('ad_soyad')},</h3><p>Sisteme girdiğiniz bildirim idare tarafından iptal edilmiştir.</p><p><b>İptal Sebebi:</b><br><i>{iptal_notu}</i></p>"
        mail_gonder(ogretmen.get("eposta"), konu, icerik)
        
    context = idareci_verilerini_getir(idareci_id, "✅ Bildirim iptal edildi ve öğretmene bilgi verildi.")
    return templates.TemplateResponse(request=request, name="idareci_panel.html", context=context)

@app.get("/form/{tur}", response_class=HTMLResponse)
async def formu_goster(request: Request, tur: str, kullanici_id: str):
    kullanici_doc = db.collection("Kullanicilar").document(kullanici_id).get()
    kullanici = kullanici_doc.to_dict()
    kullanici["id"] = kullanici_doc.id

    mekanlar_ref = db.collection("Mekanlar").get()
    mekanlar_dict = {}
    for doc in mekanlar_ref:
        m = doc.to_dict()
        bina = m.get("bina_adi")
        if bina not in mekanlar_dict: mekanlar_dict[bina] = []
        mekanlar_dict[bina].append({"id": doc.id, "isim": f"{m.get('kat')} - {m.get('alan_adi')}"})
        
    return templates.TemplateResponse(request=request, name="bildirim_formu.html", context={"tur": tur, "kullanici_id": kullanici_id, "kullanici": kullanici, "mekanlar_json": json.dumps(mekanlar_dict)})

async def upload_to_firebase(file: UploadFile, prefix: str = ""):
    if not file or not file.filename:
        return ""
    
    img_data = await file.read()
    img = Image.open(BytesIO(img_data))
    if img.mode != 'RGB': img = img.convert('RGB')
    img.thumbnail((800, 800))
    
    output_buffer = BytesIO()
    img.save(output_buffer, "JPEG", quality=60, optimize=True)
    byte_data = output_buffer.getvalue()
    
    dosya_adi = f"kbis-medya/{prefix}{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename}"
    blob = bucket.blob(dosya_adi)
    blob.upload_from_string(byte_data, content_type="image/jpeg")
    blob.make_public()
    return blob.public_url

@app.post("/bildirim_kaydet")
async def bildirim_kaydet(request: Request, tur: str = Form(...), kullanici_id: str = Form(...), mekan_id: str = Form(...), aciliyet: int = Form(...), aciklama: str = Form(...), fotograf: UploadFile = File(None)):
    try:
        foto_yolu = await upload_to_firebase(fotograf)
        
        # YENİ: İŞLEM NUMARASI OLUŞTURMA (Sayaç)
        sayac_ref = db.collection("Sayaclar").document("bildirim_sayaci")
        sayac_ref.set({tur: firestore.Increment(1)}, merge=True)
        yeni_no = sayac_ref.get().to_dict().get(tur, 1)
        islem_no = f"R-{yeni_no}" if tur == "Risk" else f"A-{yeni_no}"
        
        db.collection("Bildirimler").add({
            "islem_no": islem_no,
            "bildiren_id": kullanici_id,
            "mekan_id": mekan_id,
            "tur": tur,
            "aciliyet": aciliyet,
            "aciklama": aciklama,
            "fotograf_yolu": foto_yolu,
            "durum": "İşlem Bekliyor",
            "tarih": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        })
        
        context = ogretmen_verilerini_getir(kullanici_id, f"Harika! {tur} bildiriminiz sisteme kaydedildi.")
        
        tmy_liste = db.collection("Kullanicilar").where("rol", "==", "Teknik Müdür Yardımcısı").limit(1).get()
        mekan_doc = db.collection("Mekanlar").document(mekan_id).get()
        mekan_adi = f"{mekan_doc.to_dict().get('bina_adi')} - {mekan_doc.to_dict().get('alan_adi')}" if mekan_doc.exists else "Bilinmeyen Konum"
        kullanici_adi = context["kullanici"].get("ad_soyad", "")
        
        if tmy_liste:
            tmy = tmy_liste[0].to_dict()
            konu_tmy = f"YENİ BİLDİRİM: {tur} ({mekan_adi})"
            icerik_tmy = f"<h3>Sisteme Yeni Bir Bildirim Girildi</h3><p><b>Konum:</b> {mekan_adi}</p><p><b>Açıklama:</b> {aciklama}</p><p><b>Bildiren:</b> {kullanici_adi}</p>"
            mail_gonder(tmy.get("eposta"), konu_tmy, icerik_tmy)

        if aciliyet == 3:
            mudur_liste = db.collection("Kullanicilar").where("rol", "==", "Müdür").limit(1).get()
            if mudur_liste:
                mudur = mudur_liste[0].to_dict()
                konu_mudur = "🚨 DİKKAT: 'Çok Riskli/Acil' Bildirim Sisteme Eklendi!"
                icerik_mudur = f"<h3 style='color:red;'>Sayın Müdürüm,</h3><p>Okulda <b>'Çok Riskli/Acil'</b> seviyeli yeni bir bildirim var.</p><p><b>Konum:</b> {mekan_adi}</p><p><b>Bildiren:</b> {kullanici_adi}</p>"
                mail_gonder(mudur.get("eposta"), konu_mudur, icerik_mudur)
                    
        return templates.TemplateResponse(request=request, name="ogretmen_panel.html", context=context)
    except Exception as e:
        import traceback
        return HTMLResponse(content=f"<h1 style='color:red;'>KAYIT HATASI:</h1><pre>{traceback.format_exc()}</pre>", status_code=500)

@app.post("/is_kapat")
async def is_kapat(request: Request, bildirim_id: str = Form(...), sef_cozum_notu: str = Form(...), fotograf: UploadFile = File(None)):
    foto_yolu = await upload_to_firebase(fotograf, prefix="cozum_")
    
    db.collection("Bildirimler").document(bildirim_id).update({
        "durum": "İdare Onayı Bekliyor",
        "sef_cozum_notu": sef_cozum_notu,
        "cozum_fotografi": foto_yolu,
        "sef_kapanis_tarihi": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
    
    return HTMLResponse("<script>alert('Çözüm idareye gönderildi!'); window.location.href='/';</script>")

@app.get("/yedek_al")
async def yedek_al(request: Request):
    bildirimler = [doc.to_dict() for doc in db.collection("Bildirimler").stream()]
    kullanicilar = [doc.to_dict() for doc in db.collection("Kullanicilar").stream()]
    mekanlar = [doc.to_dict() for doc in db.collection("Mekanlar").stream()]
    
    tam_yedek = {
        "bildirimler": bildirimler,
        "kullanicilar": kullanicilar,
        "mekanlar": mekanlar
    }
    
    json_data = json.dumps(tam_yedek, default=str, ensure_ascii=False)
    
    return Response(
        content=json_data,
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=kbis_tam_yedek.json"}
    )

@app.get("/rapor_indir")
async def rapor_indir(request: Request):
    bildirimler = db.collection("Bildirimler").order_by("tarih", direction=firestore.Query.DESCENDING).stream()
    
    kullanicilar_dict = {doc.id: doc.to_dict() for doc in db.collection("Kullanicilar").stream()}
    mekanlar_dict = {doc.id: doc.to_dict() for doc in db.collection("Mekanlar").stream()}
    
    si = io.StringIO()
    cw = csv.writer(si, dialect='excel')
    cw.writerow(['Kayıt No', 'Tarih', 'Bina', 'Kat', 'Alan', 'Aciliyet', 'Durum', 'Bildiren', 'Çözüm Notu'])
    
    for doc in bildirimler:
        b = doc.to_dict()
        mekan = mekanlar_dict.get(b.get("mekan_id", ""), {})
        bildiren = kullanicilar_dict.get(b.get("bildiren_id", ""), {})
        
        cw.writerow([
            b.get("islem_no", doc.id[-5:].upper()),
            b.get("tarih", ""),
            mekan.get("bina_adi", ""),
            mekan.get("kat", ""),
            mekan.get("alan_adi", ""),
            b.get("aciliyet", ""),
            b.get("durum", ""),
            bildiren.get("ad_soyad", ""),
            b.get("sef_cozum_notu", "")
        ])
        
    output = si.getvalue().encode('utf-8-sig')
    return Response(
        content=output,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=kbis_bildirim_raporu.csv"}
    )

@app.get("/detayli_rapor", response_class=HTMLResponse)
async def detayli_rapor(request: Request, filtre: str = "1_ay", bas: str = "", bit: str = ""):
    kullanicilar_dict = {doc.id: doc.to_dict() for doc in db.collection("Kullanicilar").stream()}
    mekanlar_dict = {doc.id: doc.to_dict() for doc in db.collection("Mekanlar").stream()}
    
    if filtre == "1_ay":
        bas_tarih = (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d 00:00:00")
        aralik_metni = "Son 1 Ay"
    elif filtre == "manuel" and bas and bit:
        bas_tarih = f"{bas} 00:00:00"
        aralik_metni = f"Özel Seçim ({bas} ile {bit} Arası)"
    else:
        bas_tarih = "2000-01-01 00:00:00"
        aralik_metni = "Tüm Zamanlar"

    bildirimler_ref = db.collection("Bildirimler").order_by("tarih", direction=firestore.Query.DESCENDING).get()
    
    kapanan_isler = []
    aktif_isler = []
    sef_performans = {}
    
    for doc in bildirimler_ref:
        b = doc.to_dict()
        b["id"] = doc.id
        
        b_tarih = b.get("tarih", "")
        if b_tarih < bas_tarih:
            continue
        if filtre == "manuel" and bit and b_tarih > f"{bit} 23:59:59":
            continue
            
        mekan = mekanlar_dict.get(b.get("mekan_id", ""), {})
        b["konum"] = f"{mekan.get('bina_adi', '')} - {mekan.get('alan_adi', '')}"
        b["bildiren_kisi"] = kullanicilar_dict.get(b.get("bildiren_id", ""), {}).get("ad_soyad", "Bilinmiyor")
        b["sef"] = kullanicilar_dict.get(b.get("atanan_sef_id", ""), {}).get("ad_soyad", "Atanmadı")
        
        ac = b.get("aciliyet", 1)
        b["seviye_renk"] = {1: "#10b981", 2: "#f59e0b", 3: "#ef4444"}.get(ac, "#64748b")
        b["seviye_metni"] = {1: "Az Riskli", 2: "Riskli", 3: "Çok Riskli"}.get(ac, "Bilinmiyor") if b.get("tur") == "Risk" else {1: "Normal", 2: "Acil", 3: "Çok Acil"}.get(ac, "Bilinmiyor")
        b["acilis"] = b_tarih
        
        if b.get("durum") in ["Kapatıldı", "Yanlış Bildirim"]:
            saat_farki = 0
            if b.get("kapanis_tarihi") and b_tarih:
                try:
                    ac_time = datetime.strptime(b_tarih, "%Y-%m-%d %H:%M:%S")
                    kp_time = datetime.strptime(b.get("kapanis_tarihi"), "%Y-%m-%d %H:%M:%S")
                    saat_farki = round((kp_time - ac_time).total_seconds() / 3600, 1)
                except:
                    pass
            b["sure_saat"] = saat_farki
            kapanan_isler.append(b)
            
            sef_adi = b["sef"]
            if sef_adi != "Atanmadı" and b.get("durum") == "Kapatıldı":
                if sef_adi not in sef_performans:
                    sef_performans[sef_adi] = {"cozulen": 0, "toplam_saat": 0}
                sef_performans[sef_adi]["cozulen"] += 1
                sef_performans[sef_adi]["toplam_saat"] += saat_farki
        else:
            aktif_isler.append(b)
            
    performans_listesi = [{"sef_adi": sef, "cozulen_is": veriler["cozulen"], "ortalama_sure": round(veriler["toplam_saat"] / veriler["cozulen"], 1) if veriler["cozulen"] > 0 else 0} for sef, veriler in sef_performans.items()]
    
    ozet = {"aktif": len(aktif_isler), "cozulen": len(kapanan_isler), "toplam": len(aktif_isler) + len(kapanan_isler)}
    
    return templates.TemplateResponse(request=request, name="detayli_rapor.html", context={
        "rapor_verisi": kapanan_isler, "aktif_rapor_verisi": aktif_isler, 
        "performans_listesi": performans_listesi, "ozet": ozet, 
        "tarih": datetime.now().strftime("%d.%m.%Y %H:%M"), "aralik_metni": aralik_metni
    })